"""CSV and Excel exports of extracted data and analysis."""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TXN_COLUMNS = ["id", "date", "description", "amount", "balance", "category", "filename"]
INVOICE_COLUMNS = [
    "id", "filename", "doc_type", "status", "vendor", "customer", "document_number", "po_number",
    "issue_date", "due_date", "currency", "subtotal", "tax", "tax_rate", "discount", "shipping", "total",
    "checks_failed",
]


def _invoice_rows(documents):
    for d in documents:
        if d["doc_type"] not in ("invoice", "receipt") and not d["fields"]:
            continue
        f = d["fields"]
        row = {k: f.get(k) for k in INVOICE_COLUMNS}
        row.update(id=d["id"], filename=d["filename"], doc_type=d["doc_type"], status=d["status"])
        row["checks_failed"] = sum(1 for c in d["checks"] if c["status"] == "fail")
        yield row


def transactions_csv(transactions) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=TXN_COLUMNS, extrasaction="ignore")
    w.writeheader()
    w.writerows(transactions)
    return buf.getvalue()


def invoices_csv(documents) -> str:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=INVOICE_COLUMNS)
    w.writeheader()
    w.writerows(_invoice_rows(documents))
    return buf.getvalue()


_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MONEY = '#,##0.00;[Red]-#,##0.00'


def _sheet(wb, title, headers, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(list(row))
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        width = max([len(str(header))] + [len(str(c.value or "")) for c in ws[letter][1:200]])
        ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
        if header in money_cols:
            for c in ws[letter][1:]:
                c.number_format = _MONEY
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def workbook(transactions, documents, summary, anomalies) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = _sheet(wb, "Summary", ["Metric", "Value"], [
        ("Period start", summary["period_start"]),
        ("Period end", summary["period_end"]),
        ("Transactions", summary["count"]),
        ("Total inflow", summary["inflow"]),
        ("Total outflow", summary["outflow"]),
        ("Net cash flow", summary["net"]),
        ("Average monthly outflow", summary["avg_monthly_burn"]),
        ("Documents processed", len(documents)),
        ("Items flagged for review", len(anomalies)),
    ])
    for row in ws.iter_rows(min_row=5, max_row=8, min_col=2, max_col=2):
        for c in row:
            c.number_format = _MONEY

    _sheet(wb, "Transactions", [c.title() for c in TXN_COLUMNS],
           ([t.get(c) for c in TXN_COLUMNS] for t in transactions), money_cols=("Amount", "Balance"))
    _sheet(wb, "Expenses by Category", ["Category", "Total"], summary["expenses_by_category"], money_cols=("Total",))
    _sheet(wb, "Income by Category", ["Category", "Total"], summary["income_by_category"], money_cols=("Total",))
    _sheet(wb, "Monthly Cash Flow", ["Month", "Inflow", "Outflow", "Net"],
           ((m["month"], m["in"], m["out"], m["net"]) for m in summary["months"]),
           money_cols=("Inflow", "Outflow", "Net"))
    _sheet(wb, "Top Payees", ["Payee", "Total Paid"], summary["top_payees"], money_cols=("Total Paid",))
    _sheet(wb, "Invoices & Receipts", [c.replace("_", " ").title() for c in INVOICE_COLUMNS],
           ([r[c] for c in INVOICE_COLUMNS] for r in _invoice_rows(documents)),
           money_cols=("Subtotal", "Tax", "Discount", "Shipping", "Total"))
    _sheet(wb, "Review Flags", ["Severity", "Kind", "Detail", "Document Id", "Transaction Id"],
           ((a["severity"], a["kind"], a["message"], a.get("document_id"), a.get("transaction_id")) for a in anomalies))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
