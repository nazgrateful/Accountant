import io

from openpyxl import load_workbook


def _upload(client, samples, *names):
    files = [(io.BytesIO((samples / n).read_bytes()), n) for n in names]
    return client.post("/upload", data={"files": files}, content_type="multipart/form-data",
                       follow_redirects=True)


def test_empty_dashboard(client):
    r = client.get("/")
    assert r.status_code == 200 and b"No documents yet" in r.data


def test_full_workflow(client, samples):
    r = _upload(client, samples, "bank_statement_march_2024.pdf", "invoice_brightworks_INV-2024-0117.pdf",
                "invoice_northwind_NW-5521.pdf", "card_transactions_q1_2024.csv", "receipt_bluebottle.pdf")
    assert r.status_code == 200
    assert b"18 transactions extracted" in r.data

    for page in ("/", "/documents", "/documents?type=invoice", "/transactions", "/review", "/rules",
                 "/transactions?category=Travel&q=uber"):
        assert client.get(page).status_code == 200, page

    dash = client.get("/").data
    assert b"Software &amp; Subscriptions" in dash

    review = client.get("/review").data
    assert b"Possible duplicate payment" in review          # Adobe charged twice
    assert b"Subtotal + tax = total" in review              # Northwind invoice error

    docs = client.get("/api/documents").get_json()
    assert len(docs) == 5
    inv = next(d for d in docs if d["fields"].get("document_number") == "NW-5521")
    assert client.get(f"/documents/{inv['id']}").status_code == 200

    # Correct the total and the failed check goes away.
    form = {k: ("" if v is None else v) for k, v in inv["fields"].items() if k != "line_items"}
    form.update(total="2641.61", doc_type="invoice")
    client.post(f"/documents/{inv['id']}/fields", data=form)
    doc = next(d for d in client.get("/api/documents").get_json() if d["id"] == inv["id"])
    assert all(c["status"] != "fail" for c in doc["checks"])
    assert doc["status"] == "reviewed"

    # Manual category override survives re-applying rules.
    txns = client.get("/api/transactions").get_json()
    mystery = next(t for t in txns if t["description"] == "MYSTERY VENDOR 99")
    assert mystery["category"] == "Uncategorised"
    client.post(f"/transactions/{mystery['id']}/category", data={"category": "Consulting"})
    client.post("/rules", data={"pattern": "mystery", "category": "Other", "priority": "1"})
    client.post("/rules/apply")
    txns = client.get("/api/transactions").get_json()
    assert next(t for t in txns if t["id"] == mystery["id"])["category"] == "Consulting"

    # Exports
    csv = client.get("/export/transactions.csv")
    assert csv.status_code == 200 and csv.data.count(b"\n") == 1 + 18 + 19
    assert b"INV-2024-0117" in client.get("/export/invoices.csv").data
    wb = load_workbook(io.BytesIO(client.get("/export/report.xlsx").data))
    assert {"Summary", "Transactions", "Expenses by Category", "Monthly Cash Flow", "Invoices & Receipts",
            "Review Flags"} <= set(wb.sheetnames)

    # Original file download and delete.
    assert client.get(f"/documents/{inv['id']}/file").data[:4] == b"%PDF"
    client.post(f"/documents/{docs[-1]['id']}/delete")
    assert len(client.get("/api/documents").get_json()) == 4


def test_api_extract_does_not_store(client, samples):
    data = {"file": (io.BytesIO((samples / "receipt_bluebottle.pdf").read_bytes()), "receipt.pdf")}
    r = client.post("/api/extract", data=data, content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["fields"]["total"] == 50.75
    assert client.get("/api/documents").get_json() == []


def test_bad_upload(client):
    r = client.post("/upload", data={"files": [(io.BytesIO(b"x"), "a.exe")]},
                    content_type="multipart/form-data", follow_redirects=True)
    assert b"Unsupported file type" in r.data
    r = client.post("/api/extract", data={"file": (io.BytesIO(b"x"), "a.exe")}, content_type="multipart/form-data")
    assert r.status_code == 415
